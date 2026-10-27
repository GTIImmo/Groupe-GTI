# -*- coding: utf-8 -*-
"""CORRECTIF 1, PIECE 2 -- LES DATES JUSTES DES MANDATS DONT HEKTOR DONNE LA VEILLE.   27/10/2026

Remplit `app_mandat_date_juste` (phase2.sqlite) depuis l'export Modelo. N'ecrit
JAMAIS `app_mandat` : le run la recopie de Hektor chaque nuit, une correction n'y
tiendrait pas. Le registre lit cette table a sa fabrication
(export_app_payload.charger_dates_justes, interrupteur APP_REGISTRE_DATES_JUSTES).

═══════════════════════════════════════════════════════════════════════════════
POURQUOI
═══════════════════════════════════════════════════════════════════════════════
L'API de Hektor coupe la date UTC : un mandat du 13/06 a Paris (12/06 22:00 UTC)
lui est rendu le 12/06. Mesure du 27/10 : ~14 400 debuts et ~13 390 fins, de 2009 a
fevrier 2026, presque tous sur des biens clos ou vendus. Deux sources independantes
le disent : l'export Modelo ET le registre legal 2014-2024 (3 987 fois la meme veille).
Detail : notice/ETUDE_CORRESPONDANCES_APP_CLASSEURS_2026-10-27.md.

═══════════════════════════════════════════════════════════════════════════════
LES GARDE-FOUS
═══════════════════════════════════════════════════════════════════════════════
① UN JOUR EXACTEMENT. On ne note une date que si Hektor dit LA VEILLE de la date
   Modelo ramenee a Paris. Tout autre ecart (6 mesures le 27/10) n'est pas ce
   defaut : il va a la revue a la main, jamais ici.
② LA LIGNE DOIT ETRE DESIGNEE SANS DOUTE : numero + ref, ou numero + annonce lue
   dans la ref, ou numero porte par UNE seule ligne. Sinon : ignoree et comptee.
③ ON NOTE CE QUE HEKTOR DIT. Le registre ne remplace que si Hektor dit encore
   cette valeur : s'il corrige un jour, sa date l'emporte.
④ LA TABLE EST A COTE, JAMAIS RECONSTRUITE : on ajoute ou on met a jour, on ne
   vide pas. Retour arriere : DROP TABLE app_mandat_date_juste (ou l'interrupteur).
⑤ RIEN SANS `--ecrire`.

    python phase2/sync/poser_dates_justes_mandats.py            (a blanc)
    python phase2/sync/poser_dates_justes_mandats.py --ecrire
"""
from __future__ import annotations

import argparse
import json
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

for flux in (sys.stdout, sys.stderr):
    try:
        flux.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

from phase2.sync.importer_mandats_recherche import date_paris  # noqa: E402

PHASE2 = ROOT / "phase2" / "phase2.sqlite"
MIROIR = ROOT / "data" / "hektor.sqlite"
EXPORT = ROOT / "mandats_extraits.xlsx"
TYPES_ADMIS = ("0", "10", "6")          # le perimetre du registre (decision du 26/08)

SCHEMA = """
CREATE TABLE IF NOT EXISTS app_mandat_date_juste (
    hektor_annonce_id TEXT NOT NULL,
    numero_mandat     TEXT NOT NULL,
    champ             TEXT NOT NULL CHECK (champ IN ('debut', 'fin')),
    date_hektor       TEXT NOT NULL,
    date_juste        TEXT NOT NULL,
    source            TEXT NOT NULL,
    pose_le           TEXT NOT NULL,
    PRIMARY KEY (hektor_annonce_id, numero_mandat, champ)
)"""


def veille(texte: str | None) -> str | None:
    try:
        return (date.fromisoformat(texte[:10]) - timedelta(days=1)).isoformat()
    except (TypeError, ValueError):
        return None


def lire_lignes(con: sqlite3.Connection):
    """Les lignes du registre : (annonce, numero) -> ce que Hektor dit (version courante)."""
    ref_de = {str(a): (r or "").strip().upper() for a, r in con.execute(
        "SELECT hektor_annonce_id, no_dossier FROM hk.hektor_annonce")}
    lignes, par_numero, par_ref = {}, defaultdict(list), {}
    places = ",".join("'%s'" % t for t in TYPES_ADMIS)
    for annonce, numero, versions_json in con.execute(
            "SELECT hektor_annonce_id, numero_mandat, versions_json FROM app_mandat"
            f" WHERE COALESCE(hektor_annonce_id,'') <> '' AND COALESCE(offre_type,'') IN ({places})"):
        try:
            versions = [v for v in json.loads(versions_json or "[]") if isinstance(v, dict)]
        except Exception:
            versions = []
        if not versions:
            continue          # pas de version Hektor : rien a corriger (voir correctif 7)
        cle = (str(annonce), str(numero))
        lignes[cle] = {"debut": str(versions[0].get("debut") or "")[:10] or None,
                       "fin": str(versions[0].get("fin") or "")[:10] or None}
        par_numero[cle[1]].append(cle)
        par_ref.setdefault((cle[1], ref_de.get(cle[0], "")), cle)
    return lignes, par_numero, par_ref


def designer(numero, ref, par_numero, par_ref):
    if ref:
        if (numero, ref) in par_ref:
            return par_ref[(numero, ref)], "numero+ref"
        # la ref legale 2026 « V<agence><annonce sur 7 chiffres> » donne l'annonce
        trouve = re.match(r"^V\d{2,3}(\d{7})$", ref)
        if trouve:
            cle = (str(int(trouve.group(1))), numero)
            if cle in par_numero.get(numero, []):
                return cle, "numero+annonce"
    candidats = par_numero.get(numero, [])
    if len(candidats) == 1:
        return candidats[0], "numero seul"
    return None, "ambigu" if candidats else "absent"


def main() -> int:
    parser = argparse.ArgumentParser(description="Pose les dates justes des mandats (Hektor donne la veille).")
    parser.add_argument("--ecrire", action="store_true")
    parser.add_argument("--base", default=str(PHASE2))
    args = parser.parse_args()

    con = sqlite3.connect(args.base, timeout=120)
    con.execute("PRAGMA busy_timeout=120000")
    con.execute("ATTACH DATABASE '%s' AS hk" % MIROIR.as_posix().replace("'", "''"))
    lignes, par_numero, par_ref = lire_lignes(con)
    print(f"registre : {len(lignes)} lignes avec une version Hektor")

    from openpyxl import load_workbook
    wb = load_workbook(EXPORT, read_only=True, data_only=True)
    it = wb.worksheets[0].iter_rows(values_only=True)
    entete = [str(c).strip() if c is not None else "" for c in next(it)]
    a_poser: dict[tuple[str, str, str], tuple[str, str]] = {}
    compte, par_an, autres = Counter(), Counter(), Counter()
    for brute in it:
        x = dict(zip(entete, brute))
        numero = str(x.get("mandate_number") or "").strip()
        if not numero:
            continue
        try:
            ref = str((json.loads(x.get("linked_product_ref") or "{}") or {}).get("product_ref") or "").strip().upper()
        except Exception:
            ref = ""
        cle, comment = designer(numero, ref, par_numero, par_ref)
        if not cle:
            compte[f"ligne non designee ({comment})"] += 1
            continue
        compte[f"designee par {comment}"] += 1
        for champ, source in (("debut", x.get("date_start") or x.get("date")), ("fin", x.get("date_end"))):
            juste = date_paris(source)
            hektor = lignes[cle][champ]
            if not juste or not hektor or juste == hektor:
                continue
            if hektor == veille(juste):                       # ① un jour exactement
                a_poser[(cle[0], cle[1], champ)] = (hektor, juste)
                par_an[(champ, juste[:4])] += 1
            else:
                autres[champ] += 1
    wb.close()

    print("\nRATTACHEMENT DES LIGNES DE L'EXPORT")
    for motif, n in compte.most_common():
        print(f"   {n:>6}  {motif}")
    debuts = sum(1 for k in a_poser if k[2] == "debut")
    print(f"\nA POSER : {debuts} debuts · {len(a_poser) - debuts} fins")
    print(f"   autres ecarts (pas la veille, NON poses -> revue a la main) : {dict(autres)}")
    print("   par annee :", " ".join(f"{a}:{par_an[('debut', a)]}/{par_an[('fin', a)]}"
                                     for a in sorted({k[1] for k in par_an})), "(debuts/fins)")

    if not args.ecrire:
        print("\nA BLANC — rien n'a ete ecrit. Relancer avec --ecrire.")
        con.close()
        return 0

    con.execute(SCHEMA)
    vu = datetime.now(timezone.utc).isoformat(timespec="seconds")
    con.executemany(
        "INSERT INTO app_mandat_date_juste VALUES (?, ?, ?, ?, ?, 'modelo', ?)"
        " ON CONFLICT(hektor_annonce_id, numero_mandat, champ) DO UPDATE SET"
        "   date_hektor = excluded.date_hektor, date_juste = excluded.date_juste",
        [(a, n, c, h, j, vu) for (a, n, c), (h, j) in a_poser.items()])
    con.commit()
    total = con.execute("SELECT champ, COUNT(*) FROM app_mandat_date_juste GROUP BY 1").fetchall()
    print(f"\nECRIT. app_mandat_date_juste : {dict(total)}")
    con.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
